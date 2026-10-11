"""
Memeriksa bahwa hasil putaran uji benar-benar menempel pada baris berkas Alur.

    python3 periksa-cakupan.py <alur.xlsx> <hasil...>

Dipakai SEBELUM isi-hasil.py. Yang dicari dua hal, dan keduanya pernah terjadi:

  * hasil yang ditulis dengan nama varian yang TIDAK ADA di berkas Alur —
    hasilnya akan hilang tanpa jejak saat kolom diisi, dan laporannya terbaca
    seolah kasus itu tidak pernah dijalankan;
  * baris Alur yang belum punya hasil — supaya jumlah "Tidak diuji" pada
    laporan akhir memang disengaja, bukan kelalaian pelaksana.
"""
import json
import sys

import openpyxl

ALUR = sys.argv[1]
HASIL = sys.argv[2:]


def muat(path):
    teks = open(path, encoding='utf-8').read().strip()
    if not teks:
        return []
    if path.endswith('.jsonl'):
        return [json.loads(b) for b in teks.splitlines() if b.strip()]
    return json.loads(teks)


ws = openpyxl.load_workbook(ALUR)['Kasus Uji']
COL = {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1)}
baris = {(ws.cell(row=r, column=COL['ID Kasus']).value,
          ws.cell(row=r, column=COL['Stasiun / Varian']).value)
         for r in range(5, ws.max_row + 1)}

hasil = {}
for f in HASIL:
    for h in muat(f):
        hasil[(h['id'], h['varian'])] = h

nyasar = sorted(k for k in hasil if k not in baris)
belum = sorted(k for k in baris if k not in hasil)

print(f'baris Alur         : {len(baris)}')
print(f'hasil tercatat     : {len(hasil)}')
print(f'menempel pada baris: {len(set(hasil) & baris)}')
print()
if nyasar:
    print(f'HASIL NYASAR ({len(nyasar)}) — id/varian tidak ada di berkas Alur, hasilnya akan hilang:')
    for k in nyasar:
        print(f'  {k[0]} | {k[1]!r}')
    print()
if belum:
    print(f'BELUM ADA HASIL ({len(belum)}):')
    for k in belum:
        print(f'  {k[0]} | {k[1]!r}')
else:
    print('Setiap baris Alur sudah punya hasil.')
sys.exit(1 if nyasar else 0)
