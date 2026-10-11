"""
Memeriksa bahwa berkas Excel yang diserahkan ke klien tidak memuat jejak internal.

    python3 -P periksa-kerapian.py <berkas.xlsx> [berkas.xlsx ...]

Keluar dengan status 1 bila ada temuan, jadi bisa dipakai sebagai penjaga
sebelum menyerahkan.

Yang dicari adalah hal-hal yang MASUK AKAL ada di catatan pengembang tetapi
TIDAK boleh sampai ke pembaca di luar tim: nama kelas dan berkas sumber,
perkakas (artisan, npm, playwright, pest), nama komponen, id spesifikasi
(`screen-145`), riwayat perbaikan ("pernah ada bug", "atas permintaan user"),
dan tanggal keputusan internal. Semuanya pernah benar-benar bocor ke salah
satu berkas ini — daftarnya bukan kehati-hatian spekulatif.

Keluarannya menyebut sel persisnya (sheet!kolom+baris) supaya bisa langsung
diperbaiki, bukan hanya "ada yang kotor".
"""
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

POLA = [
    (r'(?i)hard-?cod\w*', 'istilah implementasi'),
    (r'(?i)\blocalhost\b', 'alamat lingkungan pengembang'),
    (r'(?i)\b(npm|artisan|seeder|vitest|pest|playwright)\b', 'nama perkakas'),
    (r'(?i)\bunit test\b', 'jenis uji internal'),
    (r'\b(READING_FIELDS|ListFilterBar|RecordVerificationStatus)\b', 'nama komponen/konstanta'),
    (r'\bscreen-\d+\b', 'id spesifikasi internal'),
    (r'(?i)(pernah ada bug|atas permintaan user|\bdefek\b|\bcommit\b)', 'riwayat internal'),
    (r'\bRouteAccess\b', 'nama kelas/middleware internal'),
    (r'[A-Za-z_]+::[A-Za-z_]+', 'pemanggilan kelas PHP (Kelas::metode)'),
    (r'data-testid', 'atribut pengait uji otomatis'),
    (r'\bis-(done|pending|draft|open|closed)\b', 'nama kelas CSS keadaan'),
    (r'(?i)\bmiddleware\b', 'istilah kerangka kerja'),
    (r'\bDOM\b', 'istilah teknis peramban'),
    (r'[a-z_]+\.(?:php|ts|vue|sh|blade\.php)\b', 'nama berkas sumber'),
    (r'\b(app|resources|database)/[A-Za-z0-9_/.-]+', 'jalur berkas sumber'),
    (r'\bTODO\b|\bFIXME\b|\bXXX\b', 'penanda catatan yang belum selesai'),
    (r'\b20\d\d-\d\d-\d\d\b', 'tanggal keputusan internal'),
]

# Pengecualian, per (judul kolom, sebab) — BUKAN per sel. Kolom "Tanggal"
# memang memuat tanggal putaran uji, dan kolom catatan pengamatan memang
# memuat tanggal yang TERLIHAT DI LAYAR (rentang filter "2026-09-01..",
# tanggal record). Keduanya harus dibaca klien.
#
# Dibatasi ke satu sebab saja, bukan melewati seluruh selnya: justru kolom
# catatan pengamatan itulah yang pernah membocorkan `RouteAccess`,
# `Kelas::metode`, `data-testid`, nama kelas CSS keadaan, "middleware", dan
# "DOM" — melewatkan kolomnya utuh akan mematikan penjagaan tepat di tempat
# yang paling sering bocor.
PENGECUALIAN = {
    ('Tanggal', 'tanggal keputusan internal'),
    ('Catatan / apa yang terlihat', 'tanggal keputusan internal'),
}
KOLOM_PENANDA = {'Tanggal', 'ID Kasus'}


def judul_kolom(ws):
    """Judul kolom sheet 'Kasus Uji' ada di baris 4, bukan baris 1."""
    for baris in (4, 1):
        nilai = {c: ws.cell(row=baris, column=c).value for c in range(1, ws.max_column + 1)}
        if any(v in KOLOM_PENANDA for v in nilai.values()):
            return baris, nilai
    return None, {}


kotor = 0
for path in sys.argv[1:]:
    wb = openpyxl.load_workbook(path, data_only=True)
    temuan = []
    for ws in wb.worksheets:
        baris_judul, judul = judul_kolom(ws)
        for row in ws.iter_rows():
            for sel in row:
                if not isinstance(sel.value, str):
                    continue
                kolom = judul.get(sel.column) if (
                    baris_judul is not None and sel.row > baris_judul) else None
                for pola, sebab in POLA:
                    if (kolom, sebab) in PENGECUALIAN:
                        continue
                    m = re.search(pola, sel.value)
                    if m:
                        temuan.append(
                            f'  {ws.title}!{get_column_letter(sel.column)}{sel.row}'
                            f' — {sebab}: "{m.group(0)}"')
                        break
    print(f'{path}: {"BERSIH" if not temuan else str(len(temuan)) + " sel perlu diperiksa"}')
    for t in temuan[:40]:
        print(t)
    if len(temuan) > 40:
        print(f'  … dan {len(temuan) - 40} sel lagi')
    kotor += len(temuan)

sys.exit(1 if kotor else 0)
