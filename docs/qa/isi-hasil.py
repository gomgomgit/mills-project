"""
Mengisi kolom hasil pada berkas Alur dari keluaran pelaksana uji.

    python3 isi-hasil.py <hasil.json> <sumber.xlsx> <keluaran.xlsx>

hasil.json berisi daftar {id, varian, hasil, diamati} — keluaran
e2e-web/qa-run/jalankan.spec.ts, yang menjalankan langkah kasusnya apa
adanya dan mencatat apa yang terlihat.

ATURAN YANG DIPEGANG BERKAS INI
-------------------------------
Setiap baris berkas Alur adalah jalan yang DAPAT dikerjakan — varian yang
kasusnya tidak berlaku di sana tidak dicetak sebagai baris oleh build-xlsx.py.
Jadi tidak ada kategori "tidak berlaku" di sini, dan penyebutnya seluruh baris.

Baris yang TIDAK ada di hasil.json ditandai "Tidak diuji", bukan dibiarkan
kosong dan bukan diisi Lulus. Dokumen QA-nya sendiri menyatakan: kasus yang
prasyaratnya tidak tersedia bukan kasus yang lulus. Mengisi Lulus untuk
baris yang tidak dijalankan berarti mengarang bukti uji — dan berkas ini
diserahkan ke klien.

Sebabnya ditulis spesifik per kasus, bukan satu kalimat umum, supaya pembaca
tahu APA yang menghalanginya dan siapa yang harus menyelesaikannya.
"""
import json
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

# Beberapa berkas hasil sekaligus: pelaksana web dan pelaksana mobile menulis
# berkasnya masing-masing, dan satu putaran uji adalah gabungan keduanya.
#
#     python3 isi-hasil.py <sumber.xlsx> <keluaran.xlsx> <tanggal> <hasil...>
#
# Bentuk lama (hasil lebih dulu) masih diterima supaya perintah di README yang
# sudah beredar tidak putus.
def _pisah(argv: list[str]) -> tuple[list[str], str, str, str]:
    if argv[0].endswith(('.json', '.jsonl')):
        return [argv[0]], argv[1], argv[2], (argv[3] if len(argv) > 3 else '')
    return list(argv[3:]), argv[0], argv[1], argv[2]


HASIL_BERKAS, SUMBER, KELUARAN, TANGGAL = _pisah(sys.argv[1:])

# Dicocokkan KETAT pada (id, varian). Tanpa pelapisan "berlaku untuk semua
# varian": satu kasus dijalankan pada SATU stasiun atau SATU tingkat, dan
# hasil di Depricarping bukan bukti untuk Sterilizer. Pelapisan seperti itu
# menandai baris lulus memakai bukti yang bukan miliknya, dan kesalahannya
# tidak akan terlihat oleh pembaca berkas ini.
def _muat(path: str) -> list[dict]:
    teks = open(path, encoding='utf-8').read().strip()
    if not teks:
        return []
    if path.endswith('.jsonl'):
        return [json.loads(baris) for baris in teks.splitlines() if baris.strip()]
    return json.loads(teks)


dijalankan = {}
for berkas in HASIL_BERKAS:
    for h in _muat(berkas):
        # Entri terakhir menang: satu putaran bisa menjalankan kasus yang sama
        # lebih dari sekali saat pelaksananya diperbaiki, dan yang berlaku
        # adalah pengamatan terakhir.
        dijalankan[(h['id'], h['varian'])] = h

# Mengapa sebuah kasus tidak dapat dijalankan mesin. Dicocokkan ke ISI
# kasusnya, bukan dikarang per id — supaya alasannya benar-benar berasal
# dari apa yang kasus itu minta.
ALASAN = [
    (r'(hover|arahkan kursor)',
     'Perlu penguji manusia: menuntut hover untuk membaca keterangan.'),
    (r'(ubah jam perangkat|jam perangkat)',
     'Perlu penguji manusia: menuntut mengubah jam perangkat.'),
    (r'(matikan koneksi|mode pesawat|tanpa koneksi)',
     'Perlu penguji manusia: menuntut mematikan koneksi perangkat.'),
    (r'(390px|gulir ke kanan|menggulir mendatar)',
     'Perlu penguji manusia: menuntut pemeriksaan tata letak pada lebar perangkat.'),
    (r'(tanpa gaya|berwarna|redup|kelabu dan|warna)',
     'Perlu penguji manusia: menuntut penilaian tampilan (warna, gaya, keadaan redup).'),
    (r'melebihi batas ekspor|50\.000',
     'Prasyarat tidak tersedia di basis data uji: butuh data melebihi 50.000 baris.'),
    (r'(NOMILL|kolom mill-nya kosong|tanpa mill)',
     'Prasyarat tidak tersedia di basis data uji: tidak ada akun Supervisor tanpa mill.'),
    (r'(throttling|koneksi lambat)',
     'Perlu penguji manusia: menuntut pembatasan kecepatan jaringan.'),
    (r'(buka berkas itu dengan Excel|folder unduhan|terunduh)',
     'Perlu penguji manusia: menuntut membuka berkas unduhan di Excel/LibreOffice.'),
]

MOBILE = re.compile(r'^(MB-|LP-M-)')


def alasan_untuk(cid, langkah, hasil_dih, prasyarat):
    teks = ' '.join(filter(None, (langkah, hasil_dih, prasyarat)))
    for pola, pesan in ALASAN:
        if re.search(pola, teks, re.I):
            return pesan
    if MOBILE.match(cid):
        return ('Belum dijalankan pada putaran ini: kasus aplikasi mobile, '
                'perlu perangkat atau emulator.')
    return ('Belum dijalankan pada putaran ini: perlu penguji manual '
            'mengikuti langkah pada kolom Langkah.')


wb = openpyxl.load_workbook(SUMBER)
ws = wb['Kasus Uji']
COL = {ws.cell(row=4, column=c).value: c
       for c in range(1, ws.max_column + 1)}

F_OK = Font(name='Calibri', size=10, bold=True, color='1E7B34')
F_NG = Font(name='Calibri', size=10, bold=True, color='D20000')
F_NA = Font(name='Calibri', size=10, color='6B7280')
FILL_OK = PatternFill('solid', fgColor='E6F4EA')
FILL_NG = PatternFill('solid', fgColor='FDE8E8')

hitung = {'Lulus': 0, 'Gagal': 0, 'Tidak diuji': 0}
for r in range(5, ws.max_row + 1):
    cid = ws.cell(row=r, column=COL['ID Kasus']).value
    varian = ws.cell(row=r, column=COL['Stasiun / Varian']).value
    h = dijalankan.get((cid, varian))

    if h:
        nilai, catatan = h['hasil'], h['diamati']
        penguji = 'Pelaksana QA otomatis'
        font = F_OK if nilai == 'Lulus' else F_NG
        fill = FILL_OK if nilai == 'Lulus' else FILL_NG
    else:
        nilai = 'Tidak diuji'
        catatan = alasan_untuk(
            cid,
            ws.cell(row=r, column=COL['Langkah']).value or '',
            ws.cell(row=r, column=COL['Hasil Diharapkan']).value or '',
            ws.cell(row=r, column=COL['Prasyarat']).value or '')
        penguji = ''
        font, fill = F_NA, None

    hitung[nilai] += 1
    ws.cell(row=r, column=COL['Tanggal']).value = TANGGAL if h else ''
    ws.cell(row=r, column=COL['Penguji']).value = penguji
    sel = ws.cell(row=r, column=COL['Hasil'])
    sel.value = nilai
    sel.font = font
    if fill:
        sel.fill = fill
    ws.cell(row=r, column=COL['Catatan / apa yang terlihat']).value = catatan

# ── Berkas ini harus menyatakan cakupannya SENDIRI ──────────────────────
# Pembaca yang membuka 265 baris dengan 247 "Tidak diuji" bisa menyimpulkan
# salah dua arah: bahwa sistemnya gagal, atau bahwa putaran ini sudah
# memeriksa semuanya. Keduanya tidak benar, jadi angkanya ditulis di muka —
# di judul, dan dalam satu blok di sheet Panduan.
total = sum(hitung.values())
diuji = hitung['Lulus'] + hitung['Gagal']
# Penyebutnya seluruh baris. Berkas Alur hanya memuat jalan yang dapat
# dikerjakan — varian yang kasusnya tidak berlaku di sana tidak dicetak
# sebagai baris sama sekali, jadi tidak ada yang perlu dikecualikan di sini.

ws.cell(row=1, column=1).value = (
    'Mills Smart Log — HASIL PUTARAN UJI (Alur, Tingkat 1+2)')
ws.cell(row=2, column=1).value = (
    f'{diuji} dari {total} jalan sudah dijalankan dan hasilnya terisi '
    f'({hitung["Lulus"]} Lulus, {hitung["Gagal"]} Gagal). '
    f'{hitung["Tidak diuji"]} jalan masih menunggu penguji — kolom Catatan tiap '
    f'baris menyebut apa yang menghalanginya.')

pan = wb['Panduan']
pan.cell(row=1, column=1).value = (
    'Mills Smart Log — HASIL PUTARAN UJI (Alur, Tingkat 1+2)')
pan.cell(row=2, column=1).value = (
    f'Berkas ini adalah berkas Alur yang kolom hasilnya SUDAH TERISI untuk '
    f'{diuji} dari {total} jalan. Daftar kasusnya sendiri tidak berubah.')

JUDUL = Font(name='Calibri', size=10, bold=True)
ISI = Font(name='Calibri', size=10)
FILL_INFO = PatternFill('solid', fgColor='EAF2FB')
r0 = pan.max_row + 2
blok = [
    ('CAKUPAN PUTARAN INI',
     f'{diuji} dari {total} jalan sudah dijalankan: {hitung["Lulus"]} Lulus, '
     f'{hitung["Gagal"]} Gagal. {hitung["Tidak diuji"]} jalan bertanda '
     '"Tidak diuji" — belum diperiksa, BUKAN gagal dan BUKAN lulus.'),
    ('APA YANG DIJALANKAN',
     'Langkah setiap kasus dijalankan apa adanya di peramban — layar web '
     'Laravel/Livewire dan aplikasi mobile lewat server pengembangannya — '
     'lalu APA YANG TERLIHAT dicatat di kolom Catatan. Setiap "Lulus" '
     'membawa angka dan kalimat yang dibacanya, supaya dapat ditelusuri '
     'ulang dan tidak perlu dipercaya begitu saja.'),
    ('APA YANG BELUM',
     'Yang menuntut penilaian mata atau perangkat sungguhan: warna dan '
     'keadaan redup, keterangan yang muncul saat kursor diarahkan, tata '
     'letak pada lebar ponsel sebenarnya, membuka berkas unduhan di Excel, '
     'dan perilaku aplikasi terpasang (notifikasi, izin penyimpanan). '
     'Ditambah kasus yang prasyarat datanya belum ada di basis data uji — '
     'mis. ekspor yang butuh lebih dari 50.000 baris. Alasannya ditulis '
     'pada barisnya masing-masing, bukan sebagai satu kalimat umum.'),
    ('CARA MELANJUTKAN',
     'Saring kolom Hasil ke "Tidak diuji", lalu kerjakan dari Prioritas '
     'Tinggi. Timpa nilainya saat kasusnya dijalankan. Baris "Gagal" sudah '
     'memuat apa yang terlihat; isi kolom No. Bug saat tiketnya dibuat. '
     'Baris yang sudah terisi tidak perlu diulang kecuali ada yang berubah '
     'di layarnya.'),
]
for i, (kiri, kanan) in enumerate(blok):
    r = r0 + i
    a = pan.cell(row=r, column=1)
    a.value, a.font, a.fill = kiri, JUDUL, FILL_INFO
    a.alignment = Alignment(vertical='top', wrap_text=True)
    b = pan.cell(row=r, column=2)
    b.value, b.font, b.fill = kanan, ISI, FILL_INFO
    b.alignment = Alignment(vertical='top', wrap_text=True)
    for c in (3, 4):
        pan.cell(row=r, column=c).fill = FILL_INFO
    pan.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    pan.row_dimensions[r].height = 15 * max(2, len(kanan) // 85 + 1)

wb.save(KELUARAN)
print(f'{KELUARAN}: ' + '  '.join(f'{k}={v}' for k, v in hitung.items()))
